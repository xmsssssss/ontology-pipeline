from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes
from lxml import etree

from kb.ids import normalize_header, normalize_text, stable_id
from kb.models import BlockRecord, DocumentRecord, FileRecord, SourceSpan, TableRecord
from kb.storage import Workspace


@dataclass
class ParsedFile:
    document: DocumentRecord
    blocks: list[BlockRecord] = field(default_factory=list)
    tables: list[TableRecord] = field(default_factory=list)


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    best = from_bytes(data).best()
    if best is not None:
        return str(best)
    for encoding in ("utf-8", "gb18030", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _document(file: FileRecord) -> DocumentRecord:
    return DocumentRecord(
        document_id=stable_id("doc", file.file_id, length=24),
        file_id=file.file_id,
        title=Path(file.file_name).stem,
        file_name=file.file_name,
        file_type="structured" if file.file_type == "structured" else "unstructured",
        source_path=file.source_path,
    )


def _block(
    document_id: str,
    order: int,
    block_type: str,
    text: str,
    *,
    page_no: int | None = None,
    char_start: int | None = None,
    char_end: int | None = None,
    metadata: dict[str, Any] | None = None,
) -> BlockRecord:
    normalized_lines = [
        re.sub(r"[ \t]+", " ", line).strip()
        for line in text.replace("\r\n", "\n").replace("\r", "\n").splitlines()
    ]
    normalized = "\n".join(normalized_lines).strip()
    return BlockRecord(
        block_id=stable_id(
            "block",
            document_id,
            order,
            block_type,
            normalized[:200],
            length=24,
        ),
        document_id=document_id,
        order=order,
        block_type=block_type,
        text=normalized,
        page_no=page_no,
        char_start=char_start,
        char_end=char_end,
        metadata=metadata or {},
    )


def _unique_headers(values: Iterable[Any]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values):
        base = normalize_header(value, index)
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def _row_to_record(
    document: DocumentRecord,
    table_id: str,
    row_index: int,
    headers: list[str],
    values: Iterable[Any],
    *,
    sheet_name: str | None = None,
) -> TableRecord | None:
    row = list(values)
    if not any(value not in (None, "") for value in row):
        return None
    columns = {
        header: row[index] if index < len(row) else None for index, header in enumerate(headers)
    }
    return TableRecord(
        table_record_id=stable_id(
            "table_row",
            document.document_id,
            table_id,
            row_index,
            json.dumps(columns, ensure_ascii=False, sort_keys=True, default=str),
            length=24,
        ),
        document_id=document.document_id,
        table_id=table_id,
        row_index=row_index,
        columns=columns,
        sheet_name=sheet_name,
        source_spans=[
            SourceSpan(
                row_index=row_index,
                sheet_name=sheet_name,
            )
        ],
    )


def _table_block(
    document_id: str,
    order: int,
    table_id: str,
    headers: list[str],
    values: Iterable[Any],
    row_index: int,
    *,
    page_no: int | None = None,
    sheet_name: str | None = None,
) -> BlockRecord:
    row = list(values)
    pairs = [
        f"{header}={row[index] if index < len(row) else ''}" for index, header in enumerate(headers)
    ]
    return _block(
        document_id,
        order,
        "table_row",
        "；".join(pairs),
        page_no=page_no,
        metadata={
            "table_id": table_id,
            "row_index": row_index,
            "sheet_name": sheet_name,
        },
    )


def _parse_text(file: FileRecord, text: str) -> ParsedFile:
    document = _document(file)
    parsed = ParsedFile(document=document)
    cursor = 0
    for order, match in enumerate(re.finditer(r".+?(?=\n\s*\n|\Z)", text, flags=re.DOTALL)):
        raw = match.group(0).strip()
        if not raw:
            continue
        start = text.find(raw, cursor)
        end = start + len(raw)
        cursor = end
        block_type = (
            "heading"
            if re.match(r"^(第[一二三四五六七八九十百千\d]+章|\d+(?:\.\d+)*\s)", raw)
            else "paragraph"
        )
        parsed.blocks.append(
            _block(
                document.document_id,
                order,
                block_type,
                raw,
                char_start=start,
                char_end=end,
            )
        )
    return parsed


def _parse_html(file: FileRecord, text: str) -> ParsedFile:
    document = _document(file)
    parsed = ParsedFile(document=document)
    soup = BeautifulSoup(text, "lxml")
    order = 0
    table_index = 0
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"]):
        if element.name == "table":
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                for row in element.find_all("tr")
            ]
            rows = [row for row in rows if any(cell for cell in row)]
            if not rows:
                continue
            table_id = stable_id("table", document.document_id, table_index, length=24)
            headers = _unique_headers(rows[0])
            for row_number, row in enumerate(rows[1:], start=2):
                parsed.tables.append(_row_to_record(document, table_id, row_number, headers, row))
                parsed.blocks.append(
                    _table_block(
                        document.document_id,
                        order,
                        table_id,
                        headers,
                        row,
                        row_number,
                    )
                )
                order += 1
            table_index += 1
            continue
        value = normalize_text(element.get_text(" ", strip=True))
        if not value:
            continue
        block_type = "heading" if element.name.startswith("h") else "paragraph"
        parsed.blocks.append(_block(document.document_id, order, block_type, value))
        order += 1
    parsed.blocks = [item for item in parsed.blocks if item is not None]
    parsed.tables = [item for item in parsed.tables if item is not None]
    return parsed


def _parse_docx(file: FileRecord, path: Path) -> ParsedFile:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = _document(file)
    parsed = ParsedFile(document=document)
    docx_document = Document(path)
    order = 0
    table_index = 0
    for child in docx_document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, docx_document)
            text = normalize_text(paragraph.text)
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "").lower()
            block_type = "heading" if "heading" in style or "标题" in style else "paragraph"
            parsed.blocks.append(_block(document.document_id, order, block_type, text))
            order += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, docx_document)
            rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
            rows = [row for row in rows if any(cell for cell in row)]
            if not rows:
                continue
            table_id = stable_id("table", document.document_id, table_index, length=24)
            headers = _unique_headers(rows[0])
            for row_number, row in enumerate(rows[1:], start=2):
                record = _row_to_record(document, table_id, row_number, headers, row)
                if record is not None:
                    parsed.tables.append(record)
                parsed.blocks.append(
                    _table_block(
                        document.document_id,
                        order,
                        table_id,
                        headers,
                        row,
                        row_number,
                    )
                )
                order += 1
            table_index += 1
    if not parsed.blocks:
        document.warnings.append("DOCX 未解析出文本内容。")
    return parsed


def _parse_pdf(file: FileRecord, path: Path) -> ParsedFile:
    from pypdf import PdfReader

    document = _document(file)
    parsed = ParsedFile(document=document)
    reader = PdfReader(str(path))
    order = 0
    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        paragraphs = [
            normalize_text(part) for part in re.split(r"\n\s*\n", text) if normalize_text(part)
        ]
        for paragraph in paragraphs:
            parsed.blocks.append(
                _block(
                    document.document_id,
                    order,
                    "paragraph",
                    paragraph,
                    page_no=page_index,
                )
            )
            order += 1
    if not parsed.blocks:
        document.warnings.append("PDF 未提取到文本，可能需要 OCR。")
    return parsed


def _parse_csv(file: FileRecord, path: Path) -> ParsedFile:
    document = _document(file)
    parsed = ParsedFile(document=document)
    text = decode_text(path.read_bytes())
    rows = list(csv.reader(io.StringIO(text)))
    rows = [row for row in rows if any(normalize_text(str(value)) for value in row)]
    if not rows:
        document.warnings.append("CSV 为空。")
        return parsed
    headers = _unique_headers(rows[0])
    table_id = stable_id("table", document.document_id, "csv", length=24)
    for row_index, row in enumerate(rows[1:], start=2):
        record = _row_to_record(document, table_id, row_index, headers, row)
        if record is not None:
            parsed.tables.append(record)
        parsed.blocks.append(
            _table_block(
                document.document_id,
                len(parsed.blocks),
                table_id,
                headers,
                row,
                row_index,
            )
        )
    return parsed


def _parse_xlsx(file: FileRecord, path: Path) -> ParsedFile:
    from openpyxl import load_workbook

    document = _document(file)
    parsed = ParsedFile(document=document)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [
                list(row)
                for row in sheet.iter_rows(values_only=True)
                if any(value not in (None, "") for value in row)
            ]
            if not rows:
                continue
            headers = _unique_headers(rows[0])
            table_id = stable_id("table", document.document_id, sheet.title, length=24)
            for row_index, row in enumerate(rows[1:], start=2):
                record = _row_to_record(
                    document,
                    table_id,
                    row_index,
                    headers,
                    row,
                    sheet_name=sheet.title,
                )
                if record is not None:
                    parsed.tables.append(record)
                parsed.blocks.append(
                    _table_block(
                        document.document_id,
                        len(parsed.blocks),
                        table_id,
                        headers,
                        row,
                        row_index,
                        sheet_name=sheet.title,
                    )
                )
    finally:
        workbook.close()
    if not parsed.tables:
        document.warnings.append("XLSX 未解析出数据行。")
    return parsed


def _parse_xls(file: FileRecord, path: Path) -> ParsedFile:
    import xlrd

    document = _document(file)
    parsed = ParsedFile(document=document)
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        for sheet in workbook.sheets():
            rows = [
                sheet.row_values(index)
                for index in range(sheet.nrows)
                if any(value not in (None, "") for value in sheet.row_values(index))
            ]
            if not rows:
                continue
            headers = _unique_headers(rows[0])
            table_id = stable_id("table", document.document_id, sheet.name, length=24)
            for row_index, row in enumerate(rows[1:], start=2):
                record = _row_to_record(
                    document,
                    table_id,
                    row_index,
                    headers,
                    row,
                    sheet_name=sheet.name,
                )
                if record is not None:
                    parsed.tables.append(record)
                parsed.blocks.append(
                    _table_block(
                        document.document_id,
                        len(parsed.blocks),
                        table_id,
                        headers,
                        row,
                        row_index,
                        sheet_name=sheet.name,
                    )
                )
    finally:
        workbook.release_resources()
    if not parsed.tables:
        document.warnings.append("XLS 未解析出数据行。")
    return parsed


def _json_rows(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [item if isinstance(item, dict) else {"value": item} for item in value]
    if isinstance(value, dict):
        list_values = [item for item in value.values() if isinstance(item, list)]
        if len(list_values) == 1 and all(isinstance(item, dict) for item in list_values[0]):
            return list_values[0]
        return [value]
    return [{"value": value}]


def _rows_to_table(
    document: DocumentRecord,
    rows: list[dict[str, Any]],
    *,
    table_name: str,
    sheet_name: str | None = None,
) -> tuple[list[TableRecord], list[BlockRecord]]:
    tables: list[TableRecord] = []
    blocks: list[BlockRecord] = []
    if not rows:
        return tables, blocks
    headers: list[str] = []
    for row in rows:
        for key in row:
            normalized = normalize_text(str(key)) or "value"
            if normalized not in headers:
                headers.append(normalized)
    table_id = stable_id("table", document.document_id, table_name, length=24)
    for order, (row_index, row) in enumerate(enumerate(rows, start=2)):
        values = [
            json.dumps(row.get(header), ensure_ascii=False)
            if isinstance(row.get(header), (dict, list))
            else row.get(header)
            for header in headers
        ]
        record = _row_to_record(
            document,
            table_id,
            row_index,
            headers,
            values,
            sheet_name=sheet_name,
        )
        if record is not None:
            tables.append(record)
        blocks.append(
            _table_block(
                document.document_id,
                order,
                table_id,
                headers,
                values,
                row_index,
                sheet_name=sheet_name,
            )
        )
    return tables, blocks


def _parse_json(file: FileRecord, path: Path) -> ParsedFile:
    document = _document(file)
    value = json.loads(decode_text(path.read_bytes()))
    rows = _json_rows(value)
    tables, blocks = _rows_to_table(document, rows, table_name="json")
    return ParsedFile(document=document, blocks=blocks, tables=tables)


def _element_value(element: etree._Element) -> Any:
    children = list(element)
    if not children:
        return normalize_text(element.text or "")
    result: dict[str, Any] = {}
    for child in children:
        tag = etree.QName(child).localname
        value = _element_value(child)
        if tag in result:
            existing = result[tag]
            result[tag] = existing + [value] if isinstance(existing, list) else [existing, value]
        else:
            result[tag] = value
    return result


def _parse_xml(file: FileRecord, path: Path) -> ParsedFile:
    document = _document(file)
    root = etree.fromstring(path.read_bytes())
    candidates = list(root)
    rows: list[dict[str, Any]] = []
    if candidates and all(len(list(item)) > 0 for item in candidates):
        for item in candidates:
            value = _element_value(item)
            if isinstance(value, dict):
                rows.append(value)
    if not rows:
        value = _element_value(root)
        rows = [value if isinstance(value, dict) else {"value": value}]
    tables, blocks = _rows_to_table(document, rows, table_name=etree.QName(root).localname)
    return ParsedFile(document=document, blocks=blocks, tables=tables)


def parse_file(workspace: Workspace, file: FileRecord) -> ParsedFile:
    path = workspace.root / file.stored_path
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _parse_text(file, decode_text(path.read_bytes()))
    if suffix in {".html", ".htm"}:
        return _parse_html(file, decode_text(path.read_bytes()))
    if suffix == ".docx":
        return _parse_docx(file, path)
    if suffix == ".pdf":
        return _parse_pdf(file, path)
    if suffix == ".csv":
        return _parse_csv(file, path)
    if suffix == ".xlsx":
        return _parse_xlsx(file, path)
    if suffix == ".xls":
        return _parse_xls(file, path)
    if suffix == ".json":
        return _parse_json(file, path)
    if suffix == ".xml":
        return _parse_xml(file, path)
    raise ValueError(f"没有可用的解析器: {file.file_name}")
